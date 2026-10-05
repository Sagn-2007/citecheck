import io
import logging
from typing import Optional
import pdfplumber
import docx

logger = logging.getLogger(__name__)

class DocumentExtractionError(Exception):
    """Exception raised for errors during document parsing and text extraction."""
    pass

def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_parts = []
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for i, page in enumerate(pdf.pages):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text.strip())
    except Exception as e:
        raise DocumentExtractionError(f"Failed to parse PDF file: {e}")
        
    final_text = "\n\n".join(text_parts).strip()
    if not final_text:
        raise DocumentExtractionError("No extractable text found in the PDF. It may be an image-only/scanned document.")
        
    return final_text

def extract_text_from_docx(file_bytes: bytes) -> str:
    text_parts = []
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        for para in doc.paragraphs:
            if para.text:
                text_parts.append(para.text.strip())
                
        # Also extract basic table text, reading left to right, top to bottom
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    if cell.text:
                        row_text.append(cell.text.strip())
                if row_text:
                    text_parts.append(" | ".join(row_text))
                    
    except Exception as e:
        raise DocumentExtractionError(f"Failed to parse DOCX file: {e}")
        
    final_text = "\n\n".join(text_parts).strip()
    if not final_text:
        raise DocumentExtractionError("No extractable text found in the DOCX document.")
        
    return final_text

def extract_document_text(file_bytes: bytes, filename: str, content_type: Optional[str] = None) -> str:
    if not file_bytes:
        raise DocumentExtractionError("The provided file is empty.")
        
    filename = filename.lower() if filename else ""
    content_type = content_type.lower() if content_type else ""
    
    is_pdf = filename.endswith(".pdf") or "pdf" in content_type
    is_docx = filename.endswith(".docx") or "officedocument.wordprocessingml" in content_type or "docx" in content_type
    
    if is_pdf:
        return extract_text_from_pdf(file_bytes)
    elif is_docx:
        return extract_text_from_docx(file_bytes)
    else:
        raise DocumentExtractionError("Unsupported file type. Please upload a PDF or DOCX file.")
